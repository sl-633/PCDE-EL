"""PoLyMer: PoLyInfo dictionary and polymer names curated from the literature.

The source files are not publicly available. The test set has no official
split, so lambda is selected by 5-fold CV over entries.
"""

from __future__ import annotations

import csv
import re
from collections import Counter
from pathlib import Path

from pcde_el.normalization import normalize

# Columns of the curated sheets that hold test names, by name type.
COLUMNS_COMMON = {
    "abbreviation": ["略称1", "略称2"],
    "alias": ["別名"],
    "brand_name": ["商品名"],
    "trivial_name": ["慣用名"],
}
COLUMNS_UNCOMMON = {
    "abbreviation": ["略称"],
    "trivial_name": ["慣用名"],
}
CATEGORY_ORDER = ["abbreviation", "alias", "brand_name", "trivial_name"]
# Cell values that mean "empty" in the curated sheets.
PLACEHOLDER ={"-", "‐", "ー", "―", "—", "–", "n/a", "na", "なし", "無", "－"}


def build_d0(names_csv: Path) -> list[tuple[str, str, str]]:
    """(name, pid, category) rows; Japanese names are dropped."""
    rows = []
    with open(names_csv, encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            pid = (row.get("PID") or "").strip()
            category = (row.get("name category") or "").strip().replace(" ", "_")
            name = (row.get("name") or "").strip()
            if category == "japanese_name" or not pid or not name:
                continue
            rows.append((name, pid, category))
    return rows


def _split_names(raw) -> list[str]:
    """A cell may hold several names separated by ; or ；."""
    if raw is None:
        return []
    out = []
    for part in re.split(r"[;；]", str(raw).strip()):
        part = part.strip()
        if part and part.casefold() not in PLACEHOLDER and normalize(part, 1):
            out.append(part)
    return out


def _read_sheet(
    path: Path, column_map: dict[str, list[str]]
) -> dict[str, dict[str, list[str]]]:
    """pid -> category -> names, read from the first sheet."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    rows = wb[wb.sheetnames[0]].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]

    data: dict[str, dict[str, list[str]]] = {}
    for values in rows:
        if values is None or all(v is None for v in values):
            continue
        row = {h: v for h, v in zip(header, values) if h}
        pid = str(row.get("PID") or "").strip()
        if not pid:
            continue
        entry = data.setdefault(pid, {})
        for category, columns in column_map.items():
            for col in columns:
                entry.setdefault(category, []).extend(_split_names(row.get(col)))
    return data


def build_test(
    common_xlsx: Path, uncommon_xlsx: Path, dedup: str = "exact"
) -> list[tuple[str, str, str, str]]:
    """(mention, pid, category, list) rows from both curated lists.

    dedup: "exact" drops identical rows, "normalized" keeps one row per
    (normalized mention, pid), "none" keeps everything.
    """
    sheets = {
        "common": _read_sheet(common_xlsx, COLUMNS_COMMON),
        "uncommon": _read_sheet(uncommon_xlsx, COLUMNS_UNCOMMON),
    }
    rows = [
        (name, pid, category, list_name)
        for list_name, data in sheets.items()
        for pid, cats in data.items()
        for category in CATEGORY_ORDER
        for name in cats.get(category, [])
    ]

    if dedup == "none":
        return rows
    seen, kept = set(), []
    for mention, pid, category, list_name in rows:
        key = (
            (mention, pid, category, list_name)
            if dedup == "exact"
            else (normalize(mention, 1), pid)
        )
        if key not in seen:
            seen.add(key)
            kept.append((mention, pid, category, list_name))
    return kept


def summarize(rows: list[tuple], pid_col: int = 1, cat_col: int = 2) -> str:
    cats = Counter(r[cat_col] for r in rows)
    return (
        f"{len(rows):,} rows, {len({r[pid_col] for r in rows}):,} entries, "
        f"categories {dict(cats)}"
    )
